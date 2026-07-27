import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from typing import List
from collections import Counter
from datetime import datetime
from ..models import AnalysisResult, Severity, Confidence
import logging

logger = logging.getLogger(__name__)

SEVERITY_COLORS = {
    Severity.CRITICAL:      PatternFill(start_color="FF4444", end_color="FF4444", fill_type="solid"),
    Severity.MAJOR:         PatternFill(start_color="FF8C00", end_color="FF8C00", fill_type="solid"),
    Severity.MODERATE:      PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid"),
    Severity.MINOR:         PatternFill(start_color="87CEEB", end_color="87CEEB", fill_type="solid"),
    Severity.INFORMATIONAL: PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid"),
}

AI_COLUMNS = [
    "AI Severity",
    "AI Comment",
    "AI Root Cause",
    "AI Confidence",
    "AI Evidence",
    "AI Expected Behavior",
    "AI Requires Investigation",
    "Reviewer Status",
    "Reviewer Comment",
]

class ExcelAnnotator:
    def annotate(self, report_path: str, results: List[AnalysisResult], output_path: str = None) -> str:
        """
        Add AI analysis columns to the Excel report and save to output_path.
        """
        output_path = output_path or report_path
        logger.info(f"Annotating Excel report. Output will be saved to {output_path}")
        
        wb = openpyxl.load_workbook(report_path)
        
        by_sheet = {}
        for res in results:
            by_sheet.setdefault(res.breach.sheet_name, []).append(res)
            
        for sheet_name, sheet_results in by_sheet.items():
            ws = wb[sheet_name]
            ai_start_col = ws.max_column + 1
            
            # Write headers
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="2F4F4F", end_color="2F4F4F", fill_type="solid")
            
            for i, col_name in enumerate(AI_COLUMNS):
                cell = ws.cell(row=1, column=ai_start_col + i, value=col_name)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", wrap_text=True)
                
            # Write data
            for res in sheet_results:
                row = res.breach.row_number
                
                # Severity
                cell = ws.cell(row=row, column=ai_start_col, value=res.severity.value)
                if res.severity in SEVERITY_COLORS:
                    cell.fill = SEVERITY_COLORS[res.severity]
                cell.font = Font(bold=True)
                
                ws.cell(row=row, column=ai_start_col + 1, value=res.comment)
                ws.cell(row=row, column=ai_start_col + 2, value=res.root_cause)
                ws.cell(row=row, column=ai_start_col + 3, value=res.confidence.value)
                
                evidence_text = "\n".join(f"[{ev.source}] {ev.reference}: '{ev.excerpt}'" for ev in res.evidence)
                cell_ev = ws.cell(row=row, column=ai_start_col + 4, value=evidence_text)
                cell_ev.alignment = Alignment(wrap_text=True)
                
                ws.cell(row=row, column=ai_start_col + 5, value="Yes" if res.expected_behavior else "No")
                ws.cell(row=row, column=ai_start_col + 6, value="Yes" if res.requires_investigation else "No")
                
            # Add Reviewer Status dropdown
            rev_status_col = ai_start_col + 7
            dv = DataValidation(type="list", formula1='"Approved,Rejected,Modified,Pending"', allow_blank=True)
            ws.add_data_validation(dv)
            for row_idx in range(2, ws.max_row + 1):
                dv.add(ws.cell(row=row_idx, column=rev_status_col))
                
            # Set column widths
            widths = [15, 60, 40, 12, 50, 15, 18, 15, 40]
            for i, w in enumerate(widths):
                col_letter = openpyxl.utils.get_column_letter(ai_start_col + i)
                ws.column_dimensions[col_letter].width = w

        self._add_summary_sheet(wb, results)
        wb.save(output_path)
        logger.info(f"Saved annotated report to {output_path}")
        return output_path

    def _add_summary_sheet(self, wb: openpyxl.Workbook, results: List[AnalysisResult]):
        if "AI Summary" in wb.sheetnames:
            del wb["AI Summary"]
            
        ws = wb.create_sheet("AI Summary", 0)
        
        ws.cell(row=1, column=1, value="AI Plausibility Analysis Summary").font = Font(bold=True, size=14)
        
        ws.cell(row=3, column=1, value="Analysis Date:")
        ws.cell(row=3, column=2, value=datetime.now().strftime("%Y-%m-%d %H:%M"))
        ws.cell(row=4, column=1, value="Total Breaches Analyzed:")
        ws.cell(row=4, column=2, value=len(results))
        
        ws.cell(row=6, column=1, value="Severity Breakdown").font = Font(bold=True, size=12)
        
        severity_counts = Counter(r.severity for r in results)
        for i, severity in enumerate(Severity):
            row = 7 + i
            ws.cell(row=row, column=1, value=severity.value).fill = SEVERITY_COLORS.get(severity)
            ws.cell(row=row, column=2, value=severity_counts.get(severity, 0))
            
        low_conf = [r for r in results if r.confidence == Confidence.LOW]
        ws.cell(row=13, column=1, value="Items Requiring Priority Review (Low Confidence)").font = Font(bold=True)
        ws.cell(row=14, column=1, value="Model")
        ws.cell(row=14, column=2, value="Rule")
        ws.cell(row=14, column=3, value="Uncertainty")
        for i, res in enumerate(low_conf):
            row = 15 + i
            ws.cell(row=row, column=1, value=res.breach.model_name)
            ws.cell(row=row, column=2, value=res.breach.rule_name)
            ws.cell(row=row, column=3, value=res.uncertainty_notes)
